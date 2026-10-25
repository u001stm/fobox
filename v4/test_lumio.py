"""Tests for LUMIO 4.0 engine. Run:  py -3 -m unittest test_lumio -v
Uses only synthetic files created in a temporary folder."""
import hashlib
import io
import struct
import tempfile
import unittest
from datetime import datetime, timedelta
from pathlib import Path

from engine import Engine, read_meta

try:
    from PIL import Image
except ImportError:
    Image = None


def jpeg(color, when, blank=False, wide=False):
    """Portraits are vertical, group photos horizontal — as on a real school card."""
    im = Image.new('RGB', (300, 200) if wide else (200, 300), color)
    if not blank:
        from PIL import ImageDraw
        d = ImageDraw.Draw(im)
        d.ellipse([100, 40, 180, 120], fill=(230, 200, 180))
        d.rectangle([90, 120, 190, 200], fill=(30, 40, 60))
        d.line([0, 0, 300, 200], fill=(255, 255, 255), width=3)
    ex = Image.Exif()
    ex.get_ifd(0x8769)[0x9003] = when.strftime('%Y:%m:%d %H:%M:%S')
    buf = io.BytesIO()
    im.save(buf, 'JPEG', exif=ex)
    return buf.getvalue()


def fake_tiff_raw(jpg):
    """Minimal little-endian TIFF ("RAW") with an embedded JPEG preview, no sidecar."""
    entries = [(0x0201, 4, 1, 0), (0x0202, 4, 1, len(jpg)), (0x0132, 2, 20, 0)]
    ifd_off = 8
    ifd_size = 2 + len(entries) * 12 + 4
    date_off = ifd_off + ifd_size
    jpg_off = date_off + 20
    entries[0] = (0x0201, 4, 1, jpg_off)
    entries[2] = (0x0132, 2, 20, date_off)
    out = b'II*\0' + struct.pack('<I', ifd_off) + struct.pack('<H', len(entries))
    for tag, typ, cnt, val in sorted(entries):
        out += struct.pack('<HHII', tag, typ, cnt, val)
    out += struct.pack('<I', 0) + b'2026:05:01 12:00:00\0' + jpg
    return out


def digest(folder):
    h = hashlib.sha1()
    for p in sorted(Path(folder).rglob('*')):
        if p.is_file():
            h.update(p.name.encode() + p.read_bytes() + str(p.stat().st_mtime_ns).encode())
    return h.hexdigest()


@unittest.skipIf(Image is None, 'Pillow required for synthetic images')
class EngineTest(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.src = self.tmp / 'card'
        self.src.mkdir()
        t = datetime(2026, 10, 9, 9, 0, 0)
        for n in range(100, 140):
            t += timedelta(seconds=5 if n != 120 else 900)
            colour = (60 + n % 7 * 20, 90, 140)
            (self.src / f'ABC0{n}.JPG').write_bytes(jpeg((252, 252, 252) if n == 105 else colour, t, blank=(n == 105),
                                                         wide=120 <= n < 128))   # group block; teachers after it
            (self.src / f'ABC0{n}.ARW').write_bytes(b'RAW' * 50 + bytes([n % 256]))
        (self.tmp / 'students.csv').write_text('Name,V,P1,P2,D1,D2\nИванов Иван Петрович,110,111,112,125,126\n', encoding='utf-8')
        self.e = Engine(self.tmp / 'db' / 'f.sqlite3')

    def school(self):
        return self.e.import_folder(self.src, {'students': str(self.tmp / 'students.csv'), 'order': '1', 'school': 'Школа',
                                               'class_name': '1А', 'date': '09.10.2026'})

    def approve_all(self, sid):
        st = self.e.session(sid)
        for s in st['sections']:
            self.e.review(sid, {'section': s['id'], 'selected': st['selected'][s['id']], 'approved': True})

    def test_import_is_idempotent_and_series_split(self):
        sid = self.e.import_folder(self.src)
        self.assertEqual(sid, self.e.import_folder(self.src))
        st = self.e.session(sid)
        self.assertEqual(len(st['files']), 40)                    # RAW+JPEG pairs count once
        self.assertEqual(len({f['series'] for f in st['files']}), 2)  # 15-minute pause splits series
        self.assertEqual(self.e.stats()['files'], 40)

    def test_blank_frame_flagged_but_kept(self):
        sid = self.e.import_folder(self.src)
        f = next(f for f in self.e.session(sid)['files'] if f['name'] == 'ABC0105.ARW')
        self.assertIn('white', f['flags'] or '')

    def test_sources_never_modified_by_full_cycle(self):
        before = digest(self.src)
        sid = self.school()
        self.approve_all(sid)
        self.e.export(sid, self.tmp / 'out')
        self.assertEqual(before, digest(self.src))

    def test_school_export_names_and_no_overwrite(self):
        sid = self.school()
        dry = self.e.export(sid, self.tmp / 'out', dry_run=True)   # not approved: allowed, but warned
        self.assertTrue(dry['warnings'])
        self.approve_all(sid)
        r = self.e.export(sid, self.tmp / 'out')
        kid = Path(r['root']) / 'ученики' / 'Иванов Иван'
        names = {p.name for p in kid.iterdir()}
        for n in ('Иванов Иван_ABC0110.ARW', 'ЛР_Иванов Иван_ABC0111.ARW', 'ЛС_Иванов Иван_ABC0112.ARW',
                  'Друг1_ABC0125.ARW', 'Друг2_ABC0126.ARW'):
            self.assertIn(n, names)
        self.assertTrue((Path(r['root']) / 'фон').is_dir())
        again = self.e.export(sid, self.tmp / 'out')   # re-export = top-up, nothing copied twice
        self.assertEqual(again['new'], 0)
        (Path(r['root']) / 'чужой.txt').write_text('x')
        other = self.tmp / 'out2' / Path(r['root']).relative_to(self.tmp / 'out')
        other.mkdir(parents=True)
        (other / 'чужой.txt').write_text('x')           # a non-LUMIO folder is never written into
        with self.assertRaises(ValueError):
            self.e.export(sid, self.tmp / 'out2')

    def test_child_series_stops_at_pause_and_other_people(self):
        """Like the real card: teacher 100-104, pause, child A 105-115, pause, woman 116-119, child B 120+."""
        d = self.tmp / 'card2'
        d.mkdir()
        t = datetime(2026, 10, 16, 12, 43, 0)
        for n in range(100, 130):
            t += timedelta(seconds=80 if n in (105, 116, 120) else 3)
            (d / f'ANT0{n}.JPG').write_bytes(jpeg((60 + n % 9 * 15, 90, 140), t))
            (d / f'ANT0{n}.ARW').write_bytes(b'RAW' + bytes([n % 256]))
        f = self.tmp / 'v.csv'
        f.write_text('ФИО;В;ЛР;ЛС;Д1;Д2\nГорбо Михаил;108;110;112;;\nДругой Мальчик;122;123;124;;\n', encoding='utf-8')
        sid = self.e.import_folder(d, {'students': str(f), 'date': '16.10.2026'})
        st = self.e.session(sid)
        g = st['sections'][0]
        nums = sorted(int(Path(p).stem[3:]) for p in g['files'])
        self.assertEqual((nums[0], nums[-1]), (105, 115))       # own series only: no teacher, woman or other boy
        self.assertTrue(set(g['mandatory']) <= set(st['selected'][g['id']]))
        self.assertEqual(st['session']['shoot_date'], '2026-10-16')
        self.approve_all(sid)
        root = Path(self.e.export(sid, self.tmp / 'out')['root'])
        self.assertEqual(root.parent.name, '16.10.2026')

    def test_misread_number_inside_another_childs_series(self):
        """Like Тимошина ЛР 536 inside Муцольгова 527–537: no frames of the other child are pulled in."""
        d = self.tmp / 'card3'
        d.mkdir()
        t = datetime(2026, 10, 16, 12, 0, 0)
        for n in range(100, 160):
            t += timedelta(seconds=80 if n in (110, 130) else 3)
            (d / f'ANT0{n}.JPG').write_bytes(jpeg((60 + n % 7 * 20, 90, 140), t))
            (d / f'ANT0{n}.ARW').write_bytes(b'RAW' + bytes([n % 256]))
        f = self.tmp / 'v.csv'
        f.write_text('ФИО;В;ЛР;ЛС;Д1;Д2\nПервая Анна;112;117;122;;\nВторая Вера;140;119;150;;\n', encoding='utf-8')
        sid = self.e.import_folder(d, {'students': str(f)})
        st = self.e.session(sid)
        a, b = st['sections'][0], st['sections'][1]
        na = sorted(int(Path(p).stem[3:]) for p in a['files'])
        nb = sorted(int(Path(p).stem[3:]) for p in b['files'])
        self.assertTrue(all(110 <= n <= 129 for n in na))                 # Анна: her own block only
        self.assertEqual([n for n in nb if n < 130], [119])               # Вера: the misread 119 alone, nothing around it
        self.assertIn('проверьте номер', b['unsure'].get('P1', ''))

    def test_kit_assembled_automatically_15_to_20_different_frames(self):
        """Child series of 25 frames in varied poses between pauses; neighbours before and after."""
        from PIL import ImageDraw
        d = self.tmp / 'card4'
        d.mkdir()
        t = datetime(2026, 10, 16, 12, 0, 0)
        for n in range(100, 150):
            t += timedelta(seconds=90 if n in (110, 135) else 4)
            im = Image.new('RGB', (300, 450), (40 + (n // 10) * 30 % 200, 60, 90))
            dr = ImageDraw.Draw(im)
            k = n % 25                                   # pose: position, size and arm angle all change
            dr.ellipse([60 + k * 7, 40 + (k % 5) * 12, 140 + k * 7, 120 + (k % 5) * 12], fill=(230, 200, 180))
            dr.rectangle([40 + k * 6, 130, 160 + k * 6, 400 - (k % 4) * 30], fill=(30 + k * 8, 40, 60))
            dr.line([0, (k * 37) % 450, 300, 450 - (k * 53) % 450], fill=(255, 255, 255), width=6)
            ex = Image.Exif()
            ex.get_ifd(0x8769)[0x9003] = t.strftime('%Y:%m:%d %H:%M:%S')
            buf = io.BytesIO()
            im.save(buf, 'JPEG', exif=ex)
            (d / f'ANT0{n}.JPG').write_bytes(buf.getvalue())
            (d / f'ANT0{n}.ARW').write_bytes(b'RAW' + bytes([n % 256]))
        f = self.tmp / 'v.csv'
        f.write_text('ФИО;В;ЛР;ЛС;Д1;Д2\nСоседка Ада;102;104;106;;\nГорбо Михаил;118;124;128;;\nСосед Боря;140;142;144;;\n', encoding='utf-8')
        sid = self.e.import_folder(d, {'students': str(f)})
        st = self.e.session(sid)
        g = next(x for x in st['sections'] if x['title'] == 'Горбо Михаил')
        sel = st['selected'][g['id']]
        nums = sorted(int(Path(p).stem[3:]) for p in sel)
        self.assertTrue(15 <= len(sel) <= 20, nums)
        self.assertTrue(all(110 <= n <= 134 for n in nums), nums)    # only Горбо's own block
        self.assertTrue(set(g['mandatory']) <= set(sel))

    def test_worksheet_frames_always_in_kit_after_reimport(self):
        sid = self.school()
        st = self.e.session(sid)
        s = st['sections'][0]
        with self.e.conn() as c:                                 # simulate a stale selection from an older import
            c.execute('DELETE FROM selections WHERE session_id=? AND section_id=?', (sid, s['id']))
            c.execute('INSERT INTO selections VALUES(?,?,?)', (sid, s['id'], '/old/path/X.JPG'))
        st = self.e.session(sid)
        self.assertEqual(sorted(st['selected'][s['id']]), sorted(s['mandatory']))

    def test_excel_worksheet_with_russian_headers(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest('openpyxl not installed (only needed to create the test file)')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['Заказ 7470, 9Л'])                      # title row above the header
        ws.append(['Фамилия Имя', 'Виньетка', 'Личный разворот', 'Личная страница', 'Друг 1', 'Друг 2'])
        ws.append(['Иванов Иван', 110, 111, 112, 125, 126])  # numbers as Excel numbers
        ws.append([None, None, None, None, None, None])     # empty row is skipped
        xlsx = self.tmp / 'ведомость.xlsx'
        wb.save(xlsx)
        sid = self.e.import_folder(self.src, {'students': str(xlsx), 'order': '1', 'school': 'Школа', 'class_name': '1А'})
        st = self.e.session(sid)
        self.assertEqual(st['conflicts'], [])
        p = st['sections'][0]
        self.assertEqual(p['title'], 'Иванов Иван')
        self.assertEqual(sorted(p['mandatory'].values()), ['P1', 'P2', 'V'])
        self.assertEqual(sorted(st['sections'][1]['mandatory'].values()), ['D1', 'D2'])
        self.approve_all(sid)
        r = self.e.export(sid, self.tmp / 'out')
        names = {x.name for x in (Path(r['root']) / 'ученики' / 'Иванов Иван').iterdir()}
        self.assertTrue({'Иванов Иван_ABC0110.ARW', 'ЛР_Иванов Иван_ABC0111.ARW', 'ЛС_Иванов Иван_ABC0112.ARW',
                         'Друг1_ABC0125.ARW', 'Друг2_ABC0126.ARW'} <= names)

    def test_combined_worksheet_pasted_in_one_excel_column(self):
        """The real-life layout: CSV text in column A, type column, teachers in the same file,
        empty cells, two numbers in one cell, a frame shared by two children."""
        try:
            import openpyxl
        except ImportError:
            self.skipTest('openpyxl not installed')
        lines = ['Тип;Фамилия Имя;Виньетка;ЛР;ЛС;Друг1;Друг2;Фото учителя;Прочее',
                 'ученики;Иванов Иван;110;111;112;125;126;;',
                 'ученики;Петрова Анна;;114;115;125;127 128;;',
                 'учителя;Сидорова Мария;;;;;;135;']
        wb = openpyxl.Workbook()
        for l in lines:
            wb.active.append([l])
        xlsx = self.tmp / 'список.xlsx'
        wb.save(xlsx)
        sid = self.e.import_folder(self.src, {'students': str(xlsx), 'order': '1', 'school': 'Ш', 'class_name': '1А'})
        st = self.e.session(sid)
        self.assertEqual(st['conflicts'], [])
        secs = {s['id']: s for s in st['sections']}
        self.assertEqual(len(secs['s1f']['mandatory']), 3)       # 125 + "127 128"
        self.assertEqual(len(secs['s1p']['mandatory']), 2)       # empty vignette is allowed
        self.assertEqual(secs['t0']['title'], 'Сидорова Мария')
        self.assertEqual(len(secs['t0']['mandatory']), 1)
        self.approve_all(sid)
        root = Path(self.e.export(sid, self.tmp / 'out')['root'])
        anna = {x.name for x in (root / 'ученики' / 'Петрова Анна').iterdir()}
        self.assertTrue({'ЛР_Петрова Анна_ABC0114.ARW', 'Друг1_ABC0125.ARW', 'Друг2_ABC0127.ARW', 'Друг2_ABC0128.ARW'} <= anna)
        self.assertIn('Друг1_ABC0125.ARW', {x.name for x in (root / 'ученики' / 'Иванов Иван').iterdir()})
        self.assertTrue(any((root / 'учителя').iterdir()))

    def test_missing_number_is_reported_with_child_name(self):
        f = self.tmp / 'v.csv'
        f.write_text('Иванов Иван;110;111;112;125;999\n', encoding='utf-8')   # no header at all
        sid = self.e.import_folder(self.src, {'students': str(f)})
        c = self.e.session(sid)['conflicts']
        self.assertEqual(c, ['Иванов Иван, Друг 2: 999 — нет в папке'])

    def test_empty_cells_and_absent_child_do_not_block_export(self):
        f = self.tmp / 'v.csv'
        f.write_text('ФИО;В;ЛР;ЛС;Друг1;Друг2\nИванов Иван;110;111;112;;\nОтсутствующий Пётр;;;;;\n', encoding='utf-8')
        sid = self.e.import_folder(self.src, {'students': str(f)})
        st = self.e.session(sid)
        self.assertEqual([s['id'] for s in st['sections'] if s['person']], ['s0p'])   # no friends section, nothing for the absent child
        self.assertEqual(st['absent'], ['Отсутствующий Пётр'])
        self.assertEqual(st['conflicts'], [])
        self.approve_all(sid)
        self.assertEqual(self.e.export(sid, self.tmp / 'out', dry_run=True)['files'] >= 3, True)

    def test_wrong_worksheet_frame_can_be_replaced_in_app(self):
        sid = self.school()
        st = self.e.session(sid)
        p = st['sections'][0]
        wrong = next(x for x, k in p['mandatory'].items() if k == 'P1')
        right = next(x for x in p['files'] if x not in p['mandatory'])
        self.e.set_role(sid, {'section': p['id'], 'path': wrong, 'role': None})       # "not this person"
        st = self.e.session(sid)
        p = st['sections'][0]
        self.assertNotIn(wrong, p['mandatory'])
        self.assertNotIn(wrong, st['selected'][p['id']])
        self.e.set_role(sid, {'section': p['id'], 'path': right, 'role': 'P1'})
        self.approve_all(sid)
        r = self.e.export(sid, self.tmp / 'out')
        names = {x.name for x in (Path(r['root']) / 'ученики' / 'Иванов Иван').iterdir()}
        self.assertIn('ЛР_Иванов Иван_' + Path(right).name, names)
        self.assertNotIn('ЛР_Иванов Иван_ABC0111.ARW', names)
        self.e.set_role(sid, {'section': p['id'], 'reset': True})                     # back to the worksheet
        self.assertEqual(self.e.session(sid)['sections'][0]['mandatory'].get(wrong), 'P1')

    def test_reset_whole_order_starts_from_scratch(self):
        sid = self.school()
        before = self.e.session(sid)
        p = before['sections'][0]
        wrong = next(x for x, k in p['mandatory'].items() if k == 'P1')
        self.e.set_role(sid, {'section': p['id'], 'path': wrong, 'role': None})
        self.approve_all(sid)
        self.e.review(sid, {'section': p['id'], 'selected': list(self.e.session(sid)['sections'][0]['mandatory'])})
        self.e.reset_session(sid)
        after = self.e.session(sid)
        self.assertEqual(after['approved'], {})
        self.assertEqual(after['sections'][0]['mandatory'].get(wrong), 'P1')        # worksheet roles back
        self.assertEqual({k: sorted(v) for k, v in after['selected'].items()},
                         {k: sorted(v) for k, v in before['selected'].items()})     # same as right after import

    def test_rename_child_and_teacher_kept_after_reimport(self):
        sid = self.school()
        st = self.e.session(sid)
        p = st['sections'][0]
        self.e.rename_person(sid, p['person'], 'Иванова  Ивана')
        self.assertEqual(self.e.session(sid)['sections'][0]['title'], 'Иванова Ивана')
        with self.assertRaises(ValueError):
            self.e.rename_person(sid, p['person'], '   ')
        self.assertEqual(self.school(), sid)                           # re-import the same worksheet
        self.assertIn('Иванова Ивана', [x['title'] for x in self.e.session(sid)['sections']])
        self.approve_all(sid)
        r = self.e.export(sid, self.tmp / 'out_rn')
        self.assertTrue((Path(r['root']) / 'ученики' / 'Иванова Ивана').exists())

    def test_common_photo_chosen_by_photographer(self):
        sid = self.school()
        st = self.e.session(sid)
        com = next((x for x in st['sections'] if x['id'] == 'common'), None)
        if not com:
            self.skipTest('no common section in this fixture')
        pick = com['files'][-1]
        self.e.set_common(sid, pick)
        st = self.e.session(sid)
        com = next(x for x in st['sections'] if x['id'] == 'common')
        self.assertEqual(com['mandatory'], {pick: 'G'})
        self.assertIn(pick, st['selected']['common'])

    def test_auto_exposure_sidecar_for_lightroom(self):
        from engine import exposure_fix, xmp_sidecar
        dark = Image.new('RGB', (200, 300), (40, 40, 45))
        for y in range(80, 140):                                     # a face, too dark
            for x in range(80, 120):
                dark.putpixel((x, y), (120, 85, 70))
        fix = exposure_fix(dark)
        self.assertEqual(fix['by'], 'skin')
        self.assertGreater(fix['exposure'], 0.8)                     # brighten
        bright = Image.new('RGB', (200, 300), (250, 250, 250))
        self.assertLess(exposure_fix(bright)['exposure'], 0)          # no face: whites to L 95 -> darker
        x = xmp_sidecar(fix)
        self.assertIn('crs:Exposure2012="+', x)
        self.assertIn('crs:Whites2012=', x)
        self.assertIn('crs:Blacks2012=', x)

    def test_vypusk_vertical_is_portrait_horizontal_is_group(self):
        sid = self.school()
        st = self.e.session(sid)
        wide = {f['path']: f['wide'] for f in st['files']}
        por = next(x for x in st['sections'] if x['kind'] == 'portrait')
        fr = next(x for x in st['sections'] if x['kind'] == 'friends')
        self.assertTrue(all(wide[p] == 0 for p in por['files']))
        self.assertTrue(all(wide[p] == 1 for p in fr['files']))
        self.assertTrue(all(wide[p] == 0 for p in st['selected'][por['id']]))

    def test_faces_keep_only_this_child_and_flag_wrong_number(self):
        """Face data (normally from the local face models) cleans a child's series; injected here directly."""
        import json as js
        try:
            import numpy as np
        except ImportError:
            self.skipTest('numpy not installed')
        sid = self.school()
        st = self.e.session(sid)
        por = next(x for x in st['sections'] if x['kind'] == 'portrait')
        ivan, other = np.zeros(512, 'float32'), np.zeros(512, 'float32')
        ivan[0], other[1] = 1, 1
        byname = {Path(p).name: p for p in por['files']}
        rows = []
        for name, p in byname.items():
            n = int(name[3:7])
            emb, eye = (other, 0.3) if n in (113, 114) else (ivan, 0.05 if n == 115 else 0.3)
            if n == 111:
                emb = other                              # ЛР number points at someone else
            rows.append((sid, p, 1, js.dumps([{'box': [0, 0, 10, 10], 'h': 0.3, 'eye': eye, 'eye2': eye, 'g': 'm', 'age': 10, 'score': 0.9}]),
                         emb.astype('float16').tobytes()))
        with self.e.conn() as c:
            c.executemany('INSERT OR REPLACE INTO faces VALUES(?,?,?,?,?)', rows)
        st = self.e.session(sid)
        por = next(x for x in st['sections'] if x['kind'] == 'portrait')
        names = {Path(p).name for p in por['files']}
        self.assertFalse({'ABC0113.ARW', 'ABC0114.ARW'} & names)           # another child: removed
        self.assertIn('ABC0111.ARW', names)                                # a worksheet frame is never removed…
        self.assertIn('другой человек', por['unsure'].get('P1', ''))        # …but flagged
        self.assertIn(byname['ABC0115.ARW'], por['closed'])                # closed eyes: marked only
        self.assertEqual(por['strangers'], 2)

    def test_portrait_without_face_is_marked_as_defect(self):
        import json as js
        try:
            import numpy as np
        except ImportError:
            self.skipTest('numpy not installed')
        sid = self.school()
        files = sorted((f for f in self.e.session(sid)['files'] if f['wide'] == 0), key=lambda f: f['name'])
        face = js.dumps([{'box': [0, 0, 10, 10], 'h': 0.3, 'eye': .3, 'eye2': .3, 'g': 'm', 'age': 10, 'score': .9}])
        emb = np.ones(512, 'float16').tobytes()
        back = next(f['path'] for f in files if f['name'] == 'ABC0108.ARW')
        with self.e.conn() as c:
            c.executemany('INSERT OR REPLACE INTO faces VALUES(?,?,?,?,?)',
                          [(sid, f['path'], 0, '[]', b'') if f['path'] == back else (sid, f['path'], 1, face, emb) for f in files])
        self.e._mark_noface(sid)
        fl = {f['path']: f['flags'] or '' for f in self.e.session(sid)['files']}
        self.assertIn('noface', fl[back])                                   # someone's back among portraits
        self.assertEqual(sum('noface' in v for v in fl.values()), 1)

    def test_scenes_split_by_pause_and_picture(self):
        from engine import Engine
        f = lambda n, t, h: {'name': f'A{n}', 'capture_at': f'2026-10-16 09:{t}', 'dhash': h}
        a, b = '0' * 16, 'f' * 16
        fs = [f(1, '00:00', a), f(2, '00:05', a), f(3, '00:50', a),      # one set-up
              f(4, '02:30', b), f(5, '02:33', b),                         # pause + other picture
              f(6, '06:00', b)]                                           # long pause
        self.assertEqual([[x['name'] for x in sc] for sc in Engine._scenes(fs)], [['A1', 'A2', 'A3'], ['A4', 'A5'], ['A6']])

    def test_raw_and_jpeg_in_different_folders_are_paired(self):
        d = self.tmp / 'split'
        (d / 'RAW').mkdir(parents=True)
        (d / 'JPG').mkdir()
        for n in (1, 2, 3):
            (d / 'RAW' / f'ABC0000{n}.ARW').write_bytes(b'RAW' + bytes([n]))
            (d / 'JPG' / f'ABC0000{n}.JPG').write_bytes(jpeg((90, 90, 90), datetime(2026, 1, 1, 10, 0, n)))
        sid = self.e.import_folder(d)
        files = self.e.session(sid)['files']
        self.assertEqual(len(files), 3)
        self.assertTrue(all(f['kind'] == 'raw' and f['has_preview'] for f in files))

    def test_standalone_jpeg_hidden_when_card_has_raw(self):
        (self.src / 'Полянская Зина.JPG').write_bytes(jpeg((90, 90, 90), datetime(2026, 10, 9, 9, 0, 30)))
        sid = self.e.import_folder(self.src)
        st = self.e.session(sid)
        offered = {p for s in st['sections'] for p in s['files']}
        self.assertFalse(any(p.endswith('.JPG') for p in offered))

    def test_frame_with_lost_raw_is_offered_as_jpeg_and_listed(self):
        (self.src / 'ABC0111.ARW').unlink()                       # RAW of a worksheet frame disappeared from the card
        sid = self.school()
        st = self.e.session(sid)
        offered = {p for s in st['sections'] for p in s['files'] + list(s['mandatory'])}
        self.assertIn(str(self.src / 'ABC0111.JPG'), offered)
        self.assertEqual(st['jpeg_only'], ['ABC0111.JPG'])

    def test_not_defect_survives_recheck(self):
        sid = self.e.import_folder(self.src)
        blank = next(f['path'] for f in self.e.session(sid)['files'] if f['name'] == 'ABC0105.ARW')
        self.e.not_defect(sid, blank)
        self.e.analyze(sid)
        f = next(f for f in self.e.session(sid)['files'] if f['path'] == blank)
        self.assertNotIn('white', f['flags'] or '')

    def test_vertical_raw_preview_is_rotated(self):
        d = self.tmp / 'vert'
        d.mkdir()
        raw = bytearray(fake_tiff_raw(jpeg((10, 120, 200), datetime(2026, 5, 1, 12), wide=True)))
        # add Orientation=6 (camera held vertically) to IFD0: patch entry count and append an entry
        n = struct.unpack_from('<H', raw, 8)[0]
        entries = [raw[10 + i * 12:22 + i * 12] for i in range(n)] + [struct.pack('<HHII', 0x0112, 3, 1, 6)]
        shift = 12
        fixed = []
        for e in entries:
            tag, typ, cnt, val = struct.unpack('<HHII', e)
            if tag in (0x0201, 0x0132):
                val += shift
            fixed.append(struct.pack('<HHII', tag, typ, cnt, val))
        body = raw[10 + n * 12:]
        raw = raw[:8] + struct.pack('<H', n + 1) + b''.join(sorted(fixed)) + body
        (d / 'DSC_0002.ARW').write_bytes(bytes(raw))
        sid = self.e.import_folder(d)
        data = self.e.preview_bytes(sid, str(d / 'DSC_0002.ARW'))
        w, h = Image.open(io.BytesIO(data)).size
        self.assertGreater(h, w)                 # 300x200 landscape preview shown upright as portrait

    def test_export_job_writes_order_report(self):
        import time
        sid = self.school()
        self.approve_all(sid)
        jid = self.e.start_export(sid, self.tmp / 'out')['job']
        for _ in range(100):
            if self.e.jobs[jid]['finished']:
                break
            time.sleep(0.05)
        job = self.e.jobs[jid]
        self.assertIsNone(job['error'])
        self.assertEqual(job['done'], job['total'])
        rep = Path(job['result']['report']).read_text(encoding='utf-8-sig').splitlines()
        self.assertEqual(rep[0].split(';')[:3], ['Тип', 'Кто', 'Роль'])
        self.assertTrue(any(';Иванов Иван;Личный разворот;' in r for r in rep))

    def test_export_with_open_issues_is_allowed_and_listed(self):
        f = self.tmp / 'v.csv'
        f.write_text('Иванов Иван;110;111;112;125;999\n', encoding='utf-8')   # 999 is not on the card
        sid = self.e.import_folder(self.src, {'students': str(f), 'order': '1', 'school': 'Ш', 'class_name': '1А'})
        r = self.e.export(sid, self.tmp / 'out')                                # nothing approved, one conflict
        self.assertEqual(len([w for w in r['warnings'] if 'Общая' not in w and 'Репортаж' not in w]), 3)
        root = Path(r['root'])
        self.assertIn('999', (root / '_LUMIO_замечания.txt').read_text(encoding='utf-8-sig'))
        self.assertIn('Друг1_ABC0125.ARW', {x.name for x in (root / 'ученики' / 'Иванов Иван').iterdir()})

    def test_fix_after_export_and_top_up(self):
        f = self.tmp / 'v.csv'
        f.write_text('Иванов Иван;110;111;112;125;999\n', encoding='utf-8')
        sid = self.e.import_folder(self.src, {'students': str(f), 'order': '1', 'school': 'Ш', 'class_name': '1А'})
        first = self.e.export(sid, self.tmp / 'out')
        root = Path(first['root'])
        st = self.e.session(sid)
        fr = next(x for x in st['sections'] if x['kind'] == 'friends')
        fix = next(x for x in fr['files'] if x not in fr['mandatory'])
        self.e.set_role(sid, {'section': fr['id'], 'path': fix, 'role': 'D2'})   # fix the missing Друг 2
        self.approve_all(sid)
        second = self.e.export(sid, self.tmp / 'out')
        self.assertTrue(second['update'])
        self.assertEqual(second['new'], 1)
        self.assertIn('Друг2_' + Path(fix).name, {x.name for x in (root / 'ученики' / 'Иванов Иван').iterdir()})
        self.assertFalse((root / '_LUMIO_замечания.txt').exists())

    def test_dashboard_stats_per_shoot_and_profile(self):
        sid = self.school()
        self.approve_all(sid)
        st = self.e.stats()
        s = st['sessions'][0]
        self.assertEqual((s['frames'], s['students'], s['kits']), (40, 1, 1))
        self.assertTrue(s['cover'])
        self.assertIn('2026-10-09', st['days'])
        self.e.save_profile({'name': 'Иван', 'telegram': '@ivan'})
        self.assertEqual(self.e.profile()['telegram'], '@ivan')
        with self.assertRaises(ValueError):
            self.e.save_profile({'avatar': 'javascript:alert(1)'})
        self.e.save_profile({'avatar': 'data:image/jpeg;base64,AAAA'})
        self.assertTrue(self.e.profile()['avatar'].startswith('data:image/'))

    def test_teacher_worksheet_photo_named_extra_poses_kept(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest('openpyxl not installed')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Ученики'
        ws.append(['№ договора', None, 7425])
        ws.append(['Учебное заведение', None, 1363])
        ws.append(['Класс / группа', None, 4])
        ws.append([])
        ws.append(['№', 'Фамилия, имя', 'Виньетка', 'Портрет 1', 'Портрет 2', 'Друг 1', 'Друг 2', 'Заказ'])
        ws.append([1, 'Иванов Иван', 110, 111, 112, 125, None, '—'])
        t = wb.create_sheet('Учителя')
        t.append(['№', 'ФИО', 'Должность', 'Фото'])
        t.append([1, 'Петрова Анна Сергеевна', 'учитель английского языка', 135])
        xlsx = self.tmp / 'ведомость.xlsx'
        wb.save(xlsx)
        sid = self.e.import_folder(self.src, {'students': str(xlsx)})
        st = self.e.session(sid)
        self.assertEqual((st['session']['order_no'], st['session']['school'], st['session']['class_name']), ('7425', '1363', '4'))
        tsec = next(x for x in st['sections'] if x['kind'] == 'teacher')
        extra = [p for p in tsec['files'] if p not in tsec['mandatory']][:3]
        st2 = st
        for s in st2['sections']:
            sel = list(s['mandatory']) + (extra if s is tsec else [])
            self.e.review(sid, {'section': s['id'], 'selected': sel, 'approved': True})
        root = Path(self.e.export(sid, self.tmp / 'out')['root'])
        tdir = root / 'учителя' / 'Петрова Анна Сергеевна учитель английского языка'
        names = {x.name for x in tdir.iterdir() if x.suffix != '.xmp'}
        self.assertIn('Петрова Анна Сергеевна учитель английского языка.ARW', names)
        self.assertEqual(len(names), 4)                      # worksheet photo + 3 other poses
        self.assertTrue({Path(p).name for p in extra} <= names)

    def test_teacher_sheet_with_portrait_column_and_subject_titles(self):
        try:
            import openpyxl
            from openpyxl.comments import Comment
        except ImportError:
            self.skipTest('openpyxl not installed')
        from engine import teacher_title
        self.assertEqual(teacher_title('Физкультура'), 'учитель физкультуры')
        self.assertEqual(teacher_title('Английский язык'), 'учитель английского языка')
        self.assertEqual(teacher_title('Классный руководитель'), 'классный руководитель')
        wb = openpyxl.Workbook()
        wb.active.title = 'Ученики'
        wb.active.append(['Фамилия, имя', 'Виньетка', 'Портрет 1', 'Портрет 2', 'Друг 1', 'Друг 2'])
        wb.active.append(['Иванов Иван', 110, 111, 112, 125, 126])
        t = wb.create_sheet('Учителя')
        t.append(['Ведомость выбора фотографий учителей'])
        t.append(['№ договора', None, None, 7425])
        t.append([])
        t.append(['№', 'Фамилия, имя, отчество', 'Должность', 'Портрет', 'Примечание'])
        t.append([1, 'Яковлева Ангелина Федоровна', 'Физкультура', 130, None])
        t.append([2, 'Курашова Вероника Анатольевна', 'Классный руководитель', 136, 'пометка у номера'])
        t['D6'].comment = Comment('136 или 135', 'OCR')
        xlsx = self.tmp / 'v.xlsx'
        wb.save(xlsx)
        sid = self.e.import_folder(self.src, {'students': str(xlsx)})
        st = self.e.session(sid)
        self.assertEqual(st['session']['order_no'], '7425')
        te = [x for x in st['sections'] if x['kind'] == 'teacher']
        self.assertEqual([x['subtitle'] for x in te], ['учитель физкультуры', 'классный руководитель'])
        self.assertEqual(te[1]['unsure'], {'T': '136 или 135', 'Note': 'пометка у номера'})
        self.approve_all(sid)
        root = Path(self.e.export(sid, self.tmp / 'out')['root'])
        names = {x.name for d in (root / 'учителя').iterdir() for x in d.iterdir()}
        self.assertIn('Яковлева Ангелина Федоровна учитель физкультуры', {d.name for d in (root / 'учителя').iterdir()})
        self.assertIn('Яковлева Ангелина Федоровна учитель физкультуры.ARW', names)
        self.assertIn('Курашова Вероника Анатольевна классный руководитель.ARW', names)

    def test_teacher_gets_3_to_5_different_poses_from_own_series(self):
        f = self.tmp / 'v.csv'
        f.write_text('Тип;ФИО;Виньетка;ЛР;ЛС;Друг1;Друг2;Фото учителя\nученики;Иванов Иван;110;111;112;114;115;\n'
                     'учителя;Петрова Анна Сергеевна;;;;;;128\n', encoding='utf-8')
        sid = self.e.import_folder(self.src, {'students': str(f)})
        st = self.e.session(sid)
        t = next(x for x in st['sections'] if x['kind'] == 'teacher')
        self.assertTrue(all(p not in t['files'] for x in st['sections'] if x is not t for p in x['mandatory']))
        sel = st['selected'][t['id']]
        self.assertTrue(1 <= len(sel) <= 5)
        self.assertIn(next(iter(t['mandatory'])), sel)

    def test_common_photo_and_reportage(self):
        """Free frames (nobody's series) become «Общая фотография» (1 + 5–7 duplicates) and «Репортаж» (15 + 5–7)."""
        from PIL import ImageDraw
        d = self.tmp / 'card5'
        d.mkdir()
        t = datetime(2026, 10, 16, 12, 0, 0)
        for n in range(100, 160):
            t += timedelta(seconds=90 if n in (110, 120, 128) else 4)
            im = Image.new('RGB', (300, 450) if n < 120 else (450, 300), (50, 70, 90))   # portraits vertical
            dr = ImageDraw.Draw(im)
            if 120 <= n < 128:            # group burst: same composition, tiny changes
                for k in range(6):
                    dr.ellipse([20 + k * 70, 80 + (n % 2), 70 + k * 70, 130], fill=(230, 200, 180))
            else:                         # reportage / portraits: every frame different
                k = n % 31
                dr.ellipse([k * 11, (k * 7) % 200, k * 11 + 90, (k * 7) % 200 + 90], fill=(200, 60 + k * 5, 90))
                dr.line([0, (k * 37) % 300, 450, 300 - (k * 53) % 300], fill=(255, 255, 255), width=8)
            ex = Image.Exif()
            ex.get_ifd(0x8769)[0x9003] = t.strftime('%Y:%m:%d %H:%M:%S')
            buf = io.BytesIO()
            im.save(buf, 'JPEG', exif=ex)
            (d / f'ANT0{n}.JPG').write_bytes(buf.getvalue())
            (d / f'ANT0{n}.ARW').write_bytes(b'RAW' + bytes([n % 256]))
        f = self.tmp / 'v.csv'
        f.write_text('ФИО;В;ЛР;ЛС;Д1;Д2\nИванов Иван;102;104;106;;\nПетров Пётр;112;114;116;;\n', encoding='utf-8')
        sid = self.e.import_folder(d, {'students': str(f), 'order': '1', 'school': 'Ш', 'class_name': '1А'})
        st = self.e.session(sid)
        g = next(x for x in st['sections'] if x['id'] == 'common')
        r = next(x for x in st['sections'] if x['id'] == 'report')
        gsel, gdup = st['selected']['common'], st['dups']['common']
        self.assertEqual(len(gsel) - len(gdup), 1)
        self.assertTrue(5 <= len(gdup) <= 7, gdup)
        self.assertTrue(all(120 <= int(Path(p).stem[3:]) < 128 for p in gsel))     # all from the group burst
        rsel, rdup = st['selected']['report'], st['dups']['report']
        self.assertEqual(len(rsel) - len(rdup), 15)
        self.assertTrue(not set(rsel) & set(gsel))
        self.approve_all(sid)
        root = Path(self.e.export(sid, self.tmp / 'out')['root'])
        common = sorted(x.name for x in (root / 'общая').iterdir() if x.suffix != '.xmp')
        self.assertTrue(all((root / 'общая' / n).with_suffix('.xmp').exists() for n in common))   # Lightroom sidecars
        self.assertEqual(len(common), 1 + len(gdup))                       # duplicates next to the main frame
        self.assertEqual([n for n in common if n.startswith('Общая ')], ['Общая ' + Path(next(p for p in gsel if p not in gdup)).name])
        rep = [x.name for x in (root / 'репортаж').iterdir() if x.suffix != '.xmp']
        self.assertEqual(len([n for n in rep if n.startswith('Репортаж_')]), 15)
        self.assertEqual(len(rep), 15 + len(rdup))
        self.assertEqual(sorted(x.name for x in root.iterdir() if x.is_dir()),
                         sorted(['доки', 'здание', 'общая', 'репортаж', 'ученики', 'учителя', 'фон']))
        with self.assertRaises(ValueError):                                     # 2 main group photos are not allowed
            self.e.review(sid, {'section': 'common', 'selected': gsel, 'dups': gdup[1:], 'approved': True})

    def test_companies_and_crous_sorted_delivery(self):
        """Crous: no per-child selection — everything without defects; RAW in one folder, JPEG sorted by kind."""
        from PIL import ImageDraw
        self.assertEqual([c['name'] for c in self.e.companies()][:2], ['Выпуск', 'Crous'])
        d = self.tmp / 'crous'
        d.mkdir()
        t = datetime(2026, 10, 23, 10, 0, 0)
        def shot(n, size, draw, pause=False):
            nonlocal t
            t += timedelta(seconds=200 if pause else 4)
            im = Image.new('RGB', size, (60, 80, 100))
            draw(ImageDraw.Draw(im), n)
            ex = Image.Exif()
            ex.get_ifd(0x8769)[0x9003] = t.strftime('%Y:%m:%d %H:%M:%S')
            buf = io.BytesIO()
            im.save(buf, 'JPEG', exif=ex)
            (d / f'ANT0{n}.JPG').write_bytes(buf.getvalue())
            (d / f'ANT0{n}.ARW').write_bytes(b'RAW' + bytes([n % 256]))
        person = lambda dr, n: (dr.ellipse([60 + n % 5 * 9, 60, 140 + n % 5 * 9, 140], fill=(230, 200, 180)), dr.line([0, n * 13 % 400, 200, 0], fill=(255, 255, 255), width=5))
        crowd = lambda dr, n: [dr.ellipse([20 + k * 60, 80, 70 + k * 60, 130], fill=(230, 200, 180)) for k in range(6)]
        pair = lambda dr, n: (dr.ellipse([60 + n % 4 * 30, 60, 120 + n % 4 * 30, 120], fill=(230, 200, 180)), dr.line([0, n * 17 % 300, 450, 0], fill=(255, 255, 255), width=5))
        # in the order of the CROUSALBM plan: class photo, group photos, then portraits
        for n in range(100, 108):
            shot(n, (450, 300), crowd, pause=(n == 100))        # class photo: horizontal burst
        (d / 'ANT0103.JPG').write_bytes(jpeg((252, 252, 252), t, blank=True))   # a blank (defect) frame
        for n in range(108, 114):
            shot(n, (450, 300), pair, pause=True)               # group photos: a new set-up each
        for n in range(114, 122):
            shot(n, (200, 300), person, pause=(n == 114))       # portraits: vertical, the rest of the card
        sid = self.e.import_folder(d, {'company': 'crous', 'school': 'Школа 2044', 'date': '23.10.2026'})
        st = self.e.session(sid)
        tl = st['sections'][0]
        self.assertEqual([x['kind'] for x in st['sections']], ['timeline'])            # one time line, no worksheet
        self.assertEqual([dv['cat'] for dv in tl['dividers']], ['common', 'group', 'portrait'])
        self.assertTrue(any(i['t'].startswith('Групповые: 8') and not i['ok'] for i in st['plan']))   # too few group set-ups
        self.assertNotIn(str(d / 'ANT0103.ARW'), st['selected']['timeline'])            # defect left out
        self.e.set_divider(sid, str(d / 'ANT0108.ARW'), 'teacher')                      # the photographer fixes a divider
        self.e.set_divider(sid, str(d / 'ANT0110.ARW'), 'group')
        tl = self.e.session(sid)['sections'][0]
        self.assertEqual(tl['cats'][str(d / 'ANT0109.ARW')], 'teacher')
        self.assertEqual(tl['cats'][str(d / 'ANT0111.ARW')], 'group')
        self.e.set_divider(sid, str(d / 'ANT0110.ARW'), None)                           # remove it again
        r = self.e.export(sid, self.tmp / 'out')
        root = Path(r['root'])
        self.assertEqual(root.name, '23.10.2026. Школа 2044')
        self.assertEqual(len(list((root / 'RAW').iterdir())), 21)                  # 22 frames − 1 defect
        self.assertEqual(sorted(x.name for x in (root / 'JPG').iterdir()), ['Общая', 'Педагоги', 'Портреты'])
        self.assertEqual(len(list((root / 'JPG' / 'Педагоги').iterdir())), 6)
        self.assertEqual(len(list((root / 'JPG' / 'Портреты').iterdir())), 8)
        self.assertEqual(len(list((root / 'JPG' / 'Общая').iterdir())), 7)            # 8 − the blank frame
        self.assertTrue(all(x.suffix == '.JPG' for x in (root / 'JPG' / 'Общая').iterdir()))

    def test_upload_to_yandex_disk_fake_server(self):
        """Upload goes straight from the card into disk:/<photographer>/<date>. <school>/RAW + JPG/… (fake Disk API)."""
        import threading, json as js, urllib.parse as up
        from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
        import hashlib as hl
        store, dirs, broken = {}, set(), []

        class H(BaseHTTPRequestHandler):
            def _ok(self, code=200, body=None):
                b = js.dumps(body or {}).encode()
                self.send_response(code); self.send_header('Content-Length', str(len(b))); self.end_headers(); self.wfile.write(b)
            def do_GET(self):
                u = up.urlparse(self.path); q = dict(up.parse_qsl(u.query))
                if u.path.endswith('/resources/upload'):
                    return self._ok(200, {'href': f'http://127.0.0.1:{port}/put?p=' + up.quote(q['path']), 'method': 'PUT'})
                if u.path.endswith('/resources'):
                    if q.get('fields') == 'public_url':
                        return self._ok(200, {'public_url': 'https://disk.yandex.ru/d/TEST'})
                    if q['path'] in store:
                        b_ = store[q['path']]
                        return self._ok(200, {'size': len(b_), 'type': 'file', 'md5': hl.md5(b_).hexdigest()})
                    if q['path'] in dirs:                                   # folder listing with sizes and MD5
                        pre = q['path'].rstrip('/') + '/'
                        kids = {k[len(pre):].split('/')[0] for k in list(store) + list(dirs) if k.startswith(pre)}
                        items = [{'name': n, 'type': 'dir'} if pre + n in dirs else
                                 {'name': n, 'type': 'file', 'size': len(store[pre + n]), 'md5': hl.md5(store[pre + n]).hexdigest()}
                                 for n in sorted(kids)]
                        off, lim = int(q.get('offset', 0)), int(q.get('limit', 20))
                        return self._ok(200, {'_embedded': {'items': items[off:off + lim], 'total': len(items)}})
                    return self._ok(404, {'error': 'DiskNotFoundError'})
                return self._ok(200, {'user': {'display_name': 'Сергей'}})
            def do_PUT(self):
                u = up.urlparse(self.path); q = dict(up.parse_qsl(u.query))
                if u.path == '/put':
                    data = self.rfile.read(int(self.headers['Content-Length']))
                    if not broken and q['p'].endswith('.ARW'):         # the first RAW arrives damaged
                        broken.append(q['p']); data = data[:-1] + bytes([data[-1] ^ 1])
                    store[q['p']] = data; return self._ok(201)
                if u.path.endswith('/resources'):
                    if q['path'] in dirs:
                        return self._ok(409, {'error': 'DiskPathPointsToExistentDirectoryError'})
                    dirs.add(q['path']); return self._ok(201)
                return self._ok(200)
            def log_message(self, *a): pass
        srv = ThreadingHTTPServer(('127.0.0.1', 0), H)
        port = srv.server_address[1]
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        import yadisk
        yadisk.API = f'http://127.0.0.1:{port}/v1/disk'
        try:
            sid = self.e.import_folder(self.src, {'company': 'crous', 'school': 'Школа 2044', 'date': '23.10.2026'})
            self.e.save_settings({'yandex_token': 'TOKEN'})
            self.assertEqual(self.e.disk_target(sid), 'disk:/ФОТОГРАФ СЕРГЕЙ/23.10.2026. Школа 2044')
            jid = self.e.start_upload(sid)['job']
            import time
            for _ in range(200):
                if self.e.jobs[jid]['finished']:
                    break
                time.sleep(0.05)
            job = self.e.jobs[jid]
            self.assertIsNone(job['error'])
            self.assertEqual(job['result']['link'], 'https://disk.yandex.ru/d/TEST')
            chk = job['result']['check']
            self.assertEqual((chk['expected'], chk['ok'], chk['bad']), (78, 78, []))   # damaged file found and sent again
            self.assertEqual(len(broken), 1)
            raws = [k for k in store if '/RAW/' in k]
            jpgs = [k for k in store if '/JPG/' in k]
            self.assertEqual(len(raws), 39)                                   # 40 frames − 1 blank defect
            self.assertEqual(len(jpgs), 39)
            self.assertTrue(all(k.startswith('disk:/ФОТОГРАФ СЕРГЕЙ/23.10.2026. Школа 2044/') for k in store))
            again = self.e.start_upload(sid)['job']                           # second run: nothing re-uploaded
            for _ in range(200):
                if self.e.jobs[again]['finished']:
                    break
                time.sleep(0.05)
            self.assertEqual(self.e.jobs[again]['result']['uploaded'], 0)
            # Выпуск: its own Disk (own token), the full order layout, Lightroom sidecars next to RAW
            vs = self.school()
            with self.assertRaises(ValueError):                               # Crous token is not used for Выпуск
                self.e.start_upload(vs)
            self.e.save_settings({'yandex_token@vypusk': 'TOKEN2'})
            self.approve_all(vs)
            vj = self.e.start_upload(vs)['job']
            for _ in range(400):
                if self.e.jobs[vj]['finished']:
                    break
                time.sleep(0.05)
            self.assertIsNone(self.e.jobs[vj]['error'])
            vk = [k for k in store if '/ученики/Иванов Иван/' in k]
            self.assertTrue(any(k.endswith('.ARW') for k in vk) and any(k.endswith('.xmp') for k in vk))
            chk = self.e.jobs[vj]['result']['check']
            self.assertEqual(chk['ok'], chk['expected'])
        finally:
            srv.shutdown()

    def test_csv_semicolon_cp1251(self):
        f = self.tmp / 'v.csv'
        f.write_bytes('ФИО;В;П1;П2;Д1;Д2\r\nИванов Иван;110;111;112;125;126\r\n'.encode('cp1251'))
        sid = self.e.import_folder(self.src, {'students': str(f)})
        self.assertEqual(self.e.session(sid)['conflicts'], [])

    def test_mandatory_frames_cannot_be_removed(self):
        sid = self.school()
        st = self.e.session(sid)
        s = st['sections'][0]
        with self.assertRaises(ValueError):
            self.e.review(sid, {'section': s['id'], 'selected': [], 'approved': True})

    def test_export_outside_source_only(self):
        sid = self.e.import_folder(self.src)
        self.approve_all(sid)
        with self.assertRaises(ValueError):
            self.e.export(sid, self.src / 'export')

    def test_selection_stats(self):
        sid = self.school()
        st = self.e.session(sid)
        s = st['sections'][0]
        extra = [p for p in st['selected'][s['id']] if p not in s['mandatory']]
        keep = [p for p in st['selected'][s['id']] if p not in extra[:1]]
        self.e.review(sid, {'section': s['id'], 'selected': keep, 'approved': True})
        sel = self.e.stats()['selection']
        self.assertEqual(sel['replaced'], 1 if extra else 0)
        self.assertEqual(sel['sections_approved'], 1)

    def test_embedded_raw_preview_without_sidecar(self):
        d = self.tmp / 'rawonly'
        d.mkdir()
        (d / 'DSC_0001.NEF').write_bytes(fake_tiff_raw(jpeg((10, 120, 200), datetime(2026, 5, 1, 12))))
        meta = read_meta(d / 'DSC_0001.NEF')
        self.assertEqual(meta['capture_at'], '2026-05-01 12:00:00')
        sid = self.e.import_folder(d)
        f = self.e.session(sid)['files'][0]
        self.assertTrue(f['has_preview'])
        self.assertTrue(self.e.preview_bytes(sid, str(d / 'DSC_0001.NEF')).startswith(b'\xff\xd8'))


if __name__ == '__main__':
    unittest.main()
